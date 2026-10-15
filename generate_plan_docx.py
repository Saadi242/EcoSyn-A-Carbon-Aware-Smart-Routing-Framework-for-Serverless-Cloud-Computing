import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Set standard margins (0.8 inches for plenty of room)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Palette
    COLOR_PRIMARY = RGBColor(26, 77, 46)     # Forest Green #1A4D2E
    COLOR_SECONDARY = RGBColor(79, 111, 82)  # Muted Sage Green #4F6F52
    COLOR_DARK = RGBColor(30, 41, 59)        # Slate #1E293B
    COLOR_MUTED = RGBColor(100, 116, 139)    # Gray #64748B
    HEX_HEADER_BG = "1A4D2E"
    HEX_SUBHEADER_BG = "E8EFE9"
    HEX_ALT_ROW = "F8FAF9"
    HEX_BORDER = "CBD5E1"

    # --- Header & Title ---
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("DEPARTMENT OF INFORMATION TECHNOLOGY\nTHE UNIVERSITY OF HARIPUR")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_SECONDARY

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("FINAL YEAR PROJECT (BS IT) — 30-DAY WORKLOAD EXECUTION PLAN\n")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    r_proj = p_title.add_run("ECOSYN: A Carbon-Aware Smart Routing Framework for Serverless Cloud Computing")
    r_proj.font.name = "Arial"
    r_proj.font.size = Pt(13)
    r_proj.font.bold = True
    r_proj.font.color.rgb = COLOR_DARK

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run("Project Team: Riffatullah & Muhammad Saad | Duration: 30 Days (4 Weeks) | Status: Approved (Form-2)")
    r_meta.font.name = "Arial"
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = COLOR_MUTED

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- Executive Summary Box ---
    p_box_title = doc.add_paragraph()
    r_bt = p_box_title.add_run("1. Executive Project Overview & Scope")
    r_bt.font.name = "Arial"
    r_bt.font.size = Pt(12)
    r_bt.font.bold = True
    r_bt.font.color.rgb = COLOR_PRIMARY

    p_summary = doc.add_paragraph()
    p_summary.paragraph_format.line_spacing = 1.15
    p_summary.add_run(
        "EcoSyn is an intelligent request-routing framework designed to direct serverless function invocations to "
        "cloud data center regions characterized by lower carbon intensity (gCO2eq/kWh). This document presents a "
        "rigorous, 30-day daily workload plan designed to ensure total completion of the project, including software "
        "engineering, algorithmic modeling, comparative benchmark experiments, an interactive Streamlit web dashboard, "
        "and the full 5-chapter FYP thesis report for the University of Haripur IT Department examination panel."
    )

    # --- Phase Overview Table ---
    p_phase_title = doc.add_paragraph()
    r_pt = p_phase_title.add_run("2. High-Level 30-Day Milestone Phases")
    r_pt.font.name = "Arial"
    r_pt.font.size = Pt(12)
    r_pt.font.bold = True
    r_pt.font.color.rgb = COLOR_PRIMARY

    table_phases = doc.add_table(rows=5, cols=4)
    table_phases.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_phases.autofit = False

    phase_headers = ["Phase", "Days", "Core Focus & Primary Milestone", "Lead Responsibilities"]
    col_widths = [Inches(1.2), Inches(0.9), Inches(3.2), Inches(1.5)]

    hdr_cells = table_phases.rows[0].cells
    for i, title in enumerate(phase_headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], HEX_HEADER_BG)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    phases_data = [
        ("Phase 1: Foundations & Baselines", "Days 1–7", "Config setup, cloud region profiles, carbon trace ingestion, Round-Robin and Latency-First baselines.", "Saad (Architecture & Config)\nRiffatullah (Data & DB)"),
        ("Phase 2: EcoSyn Core & Simulation", "Days 8–14", "EcoSyn multi-objective algorithm, SLA guardrails, multi-region serverless cloud simulator.", "Saad (Optimization Router)\nRiffatullah (Cloud Simulator)"),
        ("Phase 3: Benchmarks & Dashboard", "Days 15–21", "Head-to-head benchmark runs (1,000+ requests), thesis chart generation, interactive Streamlit UI.", "Saad (Streamlit Dashboard)\nRiffatullah (Benchmark Runner)"),
        ("Phase 4: Thesis Writing & Defense", "Days 22–30", "Chapters 1–5 documentation, thesis formatting, PPT presentation design, viva demo rehearsals.", "Joint (Chapters 1–5)\nJoint (Mock Defense & Slides)")
    ]

    for row_idx, data in enumerate(phases_data, start=1):
        row_cells = table_phases.rows[row_idx].cells
        for col_idx, text in enumerate(data):
            row_cells[col_idx].text = text
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=120, right=120)
            if row_idx % 2 == 1:
                set_cell_background(row_cells[col_idx], HEX_ALT_ROW)
            p = row_cells[col_idx].paragraphs[0]
            if col_idx in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(8.5)
                run.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- Daily Workload Breakdown Table ---
    p_daily_title = doc.add_paragraph()
    r_dt = p_daily_title.add_run("3. Complete 30-Day Daily Workload Schedule")
    r_dt.font.name = "Arial"
    r_dt.font.size = Pt(12)
    r_dt.font.bold = True
    r_dt.font.color.rgb = COLOR_PRIMARY

    # 30-day detailed data
    daily_plan = [
        # WEEK 1
        ("Day 1", "Foundations", "Setup development environment (Python 3.10+, venv, Git). Finalize requirements.txt and verify IDE configs.", "Saad: Git & venv setup\nRiffatullah: Tooling & dependencies", "Clean Git repo, requirements.txt, venv active"),
        ("Day 2", "Configuration", "Build global config system (config/config.yaml). Define 4 cloud regions (eu-north-1, us-east-1, ap-south-1, sa-east-1) with PUE, coords, and prices.", "Saad: Config schema design\nRiffatullah: Regional pricing & PUE research", "config/config.yaml created and loaded via PyYAML"),
        ("Day 3", "Data Ingestion", "Build Carbon Tracker base class (BaseCarbonTracker) and offline trace loader (TraceCarbonTracker) to parse 24-hr regional grid data.", "Saad: Base class architecture\nRiffatullah: CSV dataset curation & parsing", "src/carbon_tracker/trace_tracker.py working"),
        ("Day 4", "Live Carbon API", "Integrate Electricity Maps / WattTime API connector with 5-minute in-memory caching to avoid rate limit locks.", "Saad: REST API connector\nRiffatullah: Fallback logic & error handling", "src/carbon_tracker/api_tracker.py functional"),
        ("Day 5", "Baseline Routers", "Implement Round-Robin Router (Baseline 1) and Greedy Latency-First Router (Baseline 2).", "Saad: Latency-First router\nRiffatullah: Round-Robin router", "src/router/round_robin.py and latency_first.py"),
        ("Day 6", "Persistence Layer", "Design SQLite / MySQL database schema using SQLAlchemy. Create models for invocation_logs and regional_stats.", "Saad: Database manager class\nRiffatullah: ORM tables & migrations", "src/storage/db_manager.py & models.py tested"),
        ("Day 7", "Review & Sync", "Write unit tests for carbon fetching and baseline routers. Conduct Week 1 supervisor check-in with demo.", "Joint: Testing & weekly progress log", "All Week 1 tests passing (pytest tests/)"),
        
        # WEEK 2
        ("Day 8", "Algorithm Design", "Formulate EcoSyn mathematical multi-objective scoring function: Score(r) = w_c*C_r + w_l*L_r + w_p*P_r.", "Saad: Normalization equations\nRiffatullah: Weight parameter validation", "Mathematical specification documented in docs/"),
        ("Day 9", "SLA Guardrails", "Implement Stage 1 SLA filtering logic in EcoSynRouter. Disqualify regions where latency > SLA_max. Fallback safety handler.", "Saad: Core SLA filter logic\nRiffatullah: Edge-case tests (heavy loads)", "EcoSyn SLA constraint filtering fully working"),
        ("Day 10", "EcoSyn Router", "Implement full EcoSynRouter class combining SLA filtering and Pareto multi-objective scoring.", "Saad: Core router implementation\nRiffatullah: Dynamic weighting profiles", "src/router/ecosyn_router.py completed"),
        ("Day 11", "Cloud Simulation", "Build Serverless Cloud Simulator: model function instances, cold starts (100-300ms), warm executions, and concurrency limits.", "Saad: Multi-region simulator\nRiffatullah: Function execution lifecycle", "src/simulator/cloud_environment.py operational"),
        ("Day 12", "Energy & Emissions", "Implement energy consumption model (PUE, active wattage, runtime) and carbon emission calculator (gCO2eq).", "Saad: Carbon calculator formulas\nRiffatullah: Energy estimation verification", "src/metrics/carbon_calculator.py functional"),
        ("Day 13", "Workload Generator", "Build synthetic workload trace generator modeling Poisson arrival patterns and bursty traffic spikes.", "Saad: Arrival distribution engine\nRiffatullah: Real-world Azure trace synthesis", "data/workload_traces/sample_workload.csv generated"),
        ("Day 14", "Integration Test", "Run 500 simulated function invocations through all 3 routers. Verify zero crashes, DB logging, and metric accuracy.", "Joint: End-to-end integration test", "500-request pipeline validated with telemetry"),

        # WEEK 3
        ("Day 15", "Benchmark Suite", "Develop automated benchmark runner (benchmarks/run_comparison.py) executing identical workloads across all routers.", "Saad: Multi-policy execution runner\nRiffatullah: CLI argument parser & timing", "benchmarks/run_comparison.py operational"),
        ("Day 16", "Evaluation Metrics", "Implement automated summary metrics: Carbon Reduction %, Average Response Time, Cost ($), and SLA Breach Rate.", "Saad: Aggregator statistics\nRiffatullah: Result export to CSV/JSON", "src/metrics/metrics_aggregator.py completed"),
        ("Day 17", "Thesis Plots", "Write automated plotting script (benchmarks/generate_charts.py) using Matplotlib/Seaborn. Generate 4 publication figures.", "Saad: CDF & Pareto scatter plots\nRiffatullah: Carbon & regional distribution bars", "High-res figures saved in reports/figures/"),
        ("Day 18", "Streamlit UI Setup", "Initialize Streamlit web dashboard. Build sidebar controls (weights, SLA sliders, traffic size, region toggles).", "Saad: Dashboard layout & controls\nRiffatullah: State management & config binding", "src/dashboard/app.py skeleton running"),
        ("Day 19", "World Map & Heatmap", "Embed interactive world map and regional carbon intensity color heatmaps (Plotly geo scatter/choropleth).", "Saad: Interactive map rendering\nRiffatullah: Live carbon status cards", "Real-time global carbon visualizer completed"),
        ("Day 20", "Live Dispatch Demo", "Implement real-time request dispatch animation showing invocations being routed to green regions live.", "Saad: Dynamic traffic dispatch visualizer\nRiffatullah: Telemetry feed & meter widgets", "Interactive live routing demonstration completed"),
        ("Day 21", "Stress Testing", "Run 5,000+ request load test. Optimize DB batch insertion and freeze core codebase for defense.", "Joint: Stress testing & code refactoring", "Rock-solid software release; code frozen"),

        # WEEK 4
        ("Day 22", "Report Chapter 1", "Draft Chapter 1: Introduction (Background, Problem Statement, Aims & Objectives, Project Scope, Organization).", "Saad: Problem statement & background\nRiffatullah: Aims, objectives & scope", "Chapter 1 completed (approx. 8–10 pages)"),
        ("Day 23", "Report Chapter 2", "Draft Chapter 2: Literature Review (Serverless cloud, green computing, existing routing heuristics, research gap).", "Saad: Serverless & routing analysis\nRiffatullah: Green cloud & carbon tools survey", "Chapter 2 completed with 20+ IEEE citations"),
        ("Day 24", "Report Chapter 3", "Draft Chapter 3: System Design & Methodology (Architecture diagrams, sequence diagrams, algorithm pseudocode, math models).", "Saad: Algorithm pseudocode & math\nRiffatullah: Flowcharts & DB ER diagrams", "Chapter 3 completed with high-res diagrams"),
        ("Day 25", "Report Chapter 4", "Draft Chapter 4: Implementation & Experimental Results (Setup, benchmarking analysis, graphs, discussion of findings).", "Saad: Performance graphs & Pareto analysis\nRiffatullah: Experimental setup & parameters", "Chapter 4 completed with benchmark charts"),
        ("Day 26", "Report Chapter 5", "Draft Chapter 5: Conclusion & Future Work. Compile References in IEEE format and Appendices.", "Saad: Conclusion & future directions\nRiffatullah: Reference compilation & formatting", "Complete first draft of FYP Thesis Report"),
        ("Day 27", "Thesis Review", "Comprehensive proofreading, plagiarism check, formatting check per University of Haripur guidelines, supervisor review.", "Joint: Revision & supervisor feedback", "Final thesis report PDF compiled"),
        ("Day 28", "Presentation Slides", "Design 18–20 slide PowerPoint presentation for FYP defense (Problem, Architecture, Results, Live Demo preview).", "Saad: System & results slides\nRiffatullah: Introduction & literature slides", "defense_presentation.pptx finalized"),
        ("Day 29", "Mock Defense Demo", "Conduct dry-run presentation and live software demonstration. Practice answering examiner questions.", "Joint: Rehearsal and Q&A drill", "Smooth 15-minute presentation & 5-minute demo"),
        ("Day 30", "Final FYP Submission", "Submit bound thesis reports to Department of IT, commit final code to GitHub, archive project.", "Joint: Final submission & department clearance", "Official FYP Completion & Ready for External Viva")
    ]

    table_daily = doc.add_table(rows=len(daily_plan) + 1, cols=5)
    table_daily.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_daily.autofit = False

    daily_headers = ["Day", "Phase", "Daily Workload & Specific Tasks", "Team Roles", "Expected Deliverables"]
    daily_widths = [Inches(0.7), Inches(1.1), Inches(2.7), Inches(1.3), Inches(1.4)]

    # Format header
    hdr_cells_d = table_daily.rows[0].cells
    for i, title in enumerate(daily_headers):
        hdr_cells_d[i].text = title
        set_cell_background(hdr_cells_d[i], HEX_HEADER_BG)
        set_cell_margins(hdr_cells_d[i], top=100, bottom=100, left=100, right=100)
        p = hdr_cells_d[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, item in enumerate(daily_plan, start=1):
        row_cells = table_daily.rows[row_idx].cells
        # Highlight week starts
        is_week_start = row_idx in [1, 8, 15, 22]
        bg_color = HEX_SUBHEADER_BG if is_week_start else (HEX_ALT_ROW if row_idx % 2 == 1 else "FFFFFF")

        for col_idx, text in enumerate(item):
            row_cells[col_idx].text = text
            set_cell_margins(row_cells[col_idx], top=70, bottom=70, left=90, right=90)
            set_cell_background(row_cells[col_idx], bg_color)
            p = row_cells[col_idx].paragraphs[0]
            if col_idx in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(8.0)
                if is_week_start and col_idx == 0:
                    run.font.bold = True
                run.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # --- Section 4: Risk Management & Contingency Plan ---
    p_risk_title = doc.add_paragraph()
    r_rt = p_risk_title.add_run("4. Risk Assessment & Contingency Management")
    r_rt.font.name = "Arial"
    r_rt.font.size = Pt(12)
    r_rt.font.bold = True
    r_rt.font.color.rgb = COLOR_PRIMARY

    table_risks = doc.add_table(rows=4, cols=4)
    table_risks.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_risks.autofit = False

    risk_headers = ["Identified Risk", "Impact", "Likelihood", "Mitigation Strategy"]
    hdr_cells_r = table_risks.rows[0].cells
    for i, title in enumerate(risk_headers):
        hdr_cells_r[i].text = title
        set_cell_background(hdr_cells_r[i], HEX_HEADER_BG)
        set_cell_margins(hdr_cells_r[i], top=100, bottom=100, left=100, right=100)
        p = hdr_cells_r[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    risks_data = [
        ("Carbon API Rate Limits or Downtime", "High", "Medium", "EcoSyn implements local fallback caching and pre-packaged 24-hr historical grid traces so experiments never fail without internet."),
        ("Latency SLA Violations under High Carbon Drift", "High", "Low", "Built-in Stage 1 SLA Guardrail automatically prunes regions exceeding SLA_max and falls back to lowest-latency region."),
        ("Delays in Report Writing & Defense Preparation", "Medium", "Medium", "Report writing starts on Day 22 using modular chapter templates already aligned with Haripur IT department rubrics.")
    ]

    for row_idx, r_item in enumerate(risks_data, start=1):
        row_cells = table_risks.rows[row_idx].cells
        for col_idx, text in enumerate(r_item):
            row_cells[col_idx].text = text
            set_cell_margins(row_cells[col_idx], top=70, bottom=70, left=90, right=90)
            if row_idx % 2 == 1:
                set_cell_background(row_cells[col_idx], HEX_ALT_ROW)
            p = row_cells[col_idx].paragraphs[0]
            if col_idx in [1, 2]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(8.0)
                run.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # --- Section 5: Department Sign-off Section ---
    p_sign_title = doc.add_paragraph()
    r_st = p_sign_title.add_run("5. Project Sign-off & Supervision Approvals")
    r_st.font.name = "Arial"
    r_st.font.size = Pt(12)
    r_st.font.bold = True
    r_st.font.color.rgb = COLOR_PRIMARY

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.line_spacing = 1.3
    p_sig.add_run(
        "Student 1: __________________________ (Muhammad Saad)      Date: ______________\n\n"
        "Student 2: __________________________ (Riffatullah)         Date: ______________\n\n"
        "Supervisor: __________________________ (Faculty Member)     Date: ______________\n\n"
        "FYP Coordinator / Head of Department: ___________________  Date: ______________\n"
    )

    # Save to d:\Final Year Project\plan.docx
    output_path = r"d:\Final Year Project\plan.docx"
    doc.save(output_path)
    print(f"Document successfully created and saved to: {output_path}")

if __name__ == "__main__":
    create_document()

